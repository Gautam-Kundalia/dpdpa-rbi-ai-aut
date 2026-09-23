"""
Generate DPDP_Rules_2025.docx and DPDP_Act_2023.docx from db/dpdpa.db.

Layout mirrors the actual Gazette structure: each provision is a heading
(bookmarked, for linking from Excel) followed by its verbatim clause text.
A manual, clickable table of contents sits at the top of each document.

Amendment convention: a provision is shown as amended only if
provisions.latest_change_id points at an applied change_log row with
change_origin='regulatory' and change_type != 'New Provision' — i.e. only
a change the government actually made. Our own data corrections (typos,
paraphrase fixes, restored omissions — change_origin='data_correction')
are never highlighted; they stay in change_log as an audit trail only.

Within a regulatory change, only the WORDS THAT CHANGED are highlighted,
not the whole provision: old_full_text and new_full_text are compared
word by word (see word_diff_segments below); inserted/replaced words get
a yellow highlight, removed words get a grey strikethrough shown inline
immediately before the text that replaced them. A small grey caption
("Amended by <source>, <change_id>, <date>") is added under the text.

Usage:
    python src/export_word.py
"""
from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from db import PROJECT_ROOT, get_connection
from word_diff import diff_words

DOCX_SPECS = [
    {
        "path": "docs/DPDP_Rules_2025.docx",
        "title": "The Digital Personal Data Protection Rules, 2025",
        "subtitle": "Consolidated text, generated from db/dpdpa.db — Source: Gazette Notification "
                     "G.S.R. 846(E), dated 13 November 2025",
    },
    {
        "path": "docs/DPDP_Act_2023.docx",
        "title": "The Digital Personal Data Protection Act, 2023",
        "subtitle": "Consolidated text, generated from db/dpdpa.db — Source: Act No. 22 of 2023, "
                     "as notified in phases per G.S.R. 843(E), dated 13 November 2025",
    },
]

HEADING_COLOR = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)


# --------------------------------------------------------------------------
# Low-level helpers: bookmarks and internal hyperlinks (python-docx has no
# built-in API for either, so these drop down to raw OXML).
# --------------------------------------------------------------------------

def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, text, anchor):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)

    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


# --------------------------------------------------------------------------
# Markdown-lite rendering: **bold**, blank-line paragraphs, '|' tables
# --------------------------------------------------------------------------

def add_inline_runs(paragraph, text, highlight=False, strike=False):
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 1:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if strike:
            run.font.strike = True
            run.font.color.rgb = GREY


def is_table_block(block: str) -> bool:
    first_line = block.strip().split("\n")[0]
    return first_line.strip().startswith("|")


def add_markdown_table(doc, block: str):
    lines = [l for l in block.split("\n") if l.strip()]
    rows = []
    for line in lines:
        if re.match(r"^\s*\|?\s*-+\s*(\|\s*-+\s*)*\|?\s*$", line):
            continue  # separator row
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_runs(p, cell_text)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def render_markdown_body(doc, text: str, highlight=False, strike=False):
    if not text:
        return
    blocks = text.split("\n\n")
    for block in blocks:
        if not block.strip():
            continue
        if is_table_block(block):
            add_markdown_table(doc, block)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            lines = block.split("\n")
            for i, line in enumerate(lines):
                if i > 0:
                    p.add_run().add_break()
                add_inline_runs(p, line, highlight=highlight, strike=strike)


# --------------------------------------------------------------------------
# Word-level diff: highlight only the words that changed (D1/D2)
# diff_words() lives in word_diff.py, shared with notify.py's email
# snippets, so the docx highlighting and the email always agree.
# --------------------------------------------------------------------------

def split_blocks_with_offsets(text):
    """Same paragraph split as render_markdown_body (on '\\n\\n'), but keeping
    each block's (start, end) character offset within the original text."""
    blocks = []
    offset = 0
    for part in text.split("\n\n"):
        start = offset
        end = start + len(part)
        blocks.append((start, end, part))
        offset = end + 2
    return blocks


def _find_all_spans(text, sub):
    """All non-overlapping (start, end) occurrences of `sub` in `text`."""
    spans = []
    start = 0
    while True:
        idx = text.find(sub, start)
        if idx == -1:
            break
        spans.append((idx, idx + len(sub)))
        start = idx + len(sub)
    return spans


def _add_diff_words(paragraph, old_words, new_words):
    for words, style in diff_words(old_words, new_words):
        text = " ".join(words)
        if not text:
            continue
        run = paragraph.add_run(text)
        if style == "delete":
            run.font.strike = True
            run.font.color.rgb = GREY
        elif style == "insert":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        paragraph.add_run(" ")


def render_block_with_diff(paragraph, block_text, spans, old_words, new_words):
    """
    Render one markdown-lite block into `paragraph` as a single docx
    paragraph, applying word-diff highlighting at every span in `spans`
    (a list of (start, end) character offsets within block_text) — a
    single provision can contain the SAME corrected phrase more than once
    (e.g. Rule 1(3) and 1(4) both had "of this Gazette" corrected by the
    same corrigendum item), and every occurrence must be highlighted, not
    just the first.
    """
    pos = 0
    for start, end in spans:
        prefix = block_text[pos:start]
        if prefix:
            add_inline_runs(paragraph, prefix)
        _add_diff_words(paragraph, old_words, new_words)
        pos = end
    suffix = block_text[pos:]
    if suffix:
        add_inline_runs(paragraph, suffix)


def add_amendment_caption(doc, amendment):
    date = (amendment["detected_timestamp"] or "")[:10]
    p = doc.add_paragraph()
    run = p.add_run(
        f"Amended by {amendment['source_document'] or 'official source'} "
        f"({amendment['change_id']}, {date})"
    )
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    run.italic = True


def render_provision_body(doc, full_text, amendment):
    """
    Render a provision's body text. If `amendment` is a real, applied
    regulatory change (change_origin='regulatory'), highlight only the
    changed words within full_text and add a caption — at EVERY occurrence
    of new_full_text, since one corrigendum/amendment item can correct the
    same phrase more than once within a single provision (e.g. Rule 1(3)
    and 1(4) both had "of this Gazette" corrected). Returns True if a
    warning was logged (new_full_text no longer found verbatim in
    full_text — e.g. a later data correction changed the surrounding
    text), in which case the text is still rendered, plainly, with the
    caption still shown.
    """
    full_text = full_text or ""
    if not amendment or not amendment["new_full_text"]:
        render_markdown_body(doc, full_text)
        return False

    new_full_text = amendment["new_full_text"]
    if new_full_text not in full_text:
        render_markdown_body(doc, full_text)
        add_amendment_caption(doc, amendment)
        print(
            f"WARNING: {amendment['change_id']} new_full_text is no longer a "
            f"substring of the current full_text — rendering plainly with caption only.",
            file=sys.stderr,
        )
        return True

    blocks = split_blocks_with_offsets(full_text)
    # Every block that contains at least one occurrence of new_full_text
    # (and isn't a table — those aren't rendered via this paragraph path).
    diff_blocks = {
        (start, end): _find_all_spans(block, new_full_text)
        for start, end, block in blocks
        if new_full_text in block and not is_table_block(block)
    }

    if not diff_blocks:
        render_markdown_body(doc, full_text)
        add_amendment_caption(doc, amendment)
        print(
            f"WARNING: {amendment['change_id']}'s changed text only appears inside a table "
            f"or spans multiple paragraphs — rendering plainly with caption only.",
            file=sys.stderr,
        )
        return True

    old_words = re.findall(r"\S+", amendment["old_full_text"] or "")
    new_words = re.findall(r"\S+", new_full_text)

    for start, end, block in blocks:
        spans = diff_blocks.get((start, end))
        if spans:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            render_block_with_diff(p, block, spans, old_words, new_words)
        else:
            render_markdown_body(doc, block)
    add_amendment_caption(doc, amendment)
    return False


# --------------------------------------------------------------------------
# Document assembly
# --------------------------------------------------------------------------

def fetch_provisions(conn, docx_relpath):
    rows = conn.execute(
        "SELECT provision_id, reference, topic_category, full_text, full_text_anchor, "
        "status, effective_date, notes, latest_change_id "
        "FROM provisions WHERE full_text_path = ? ORDER BY sort_order",
        (docx_relpath,),
    ).fetchall()
    return rows


def fetch_latest_amendment(conn, provision_id):
    """
    Returns the change that provisions.latest_change_id points at IF it
    represents a real, applied REGULATORY change — change_origin='regulatory'
    (the government actually changed the law) and change_type != 'New
    Provision'. change_origin='data_correction' rows (our own typo/paraphrase
    fixes) never match here, however recent they are, so they never affect
    rendering (D1). Keying on latest_change_id — rather than "the newest
    change_log row of any qualifying type" — is also what lets those
    corrections stay in the audit trail without disturbing which change (if
    any) is shown as an amendment.
    """
    row = conn.execute(
        "SELECT c.change_id, c.change_type, c.old_full_text, c.new_full_text, "
        "c.source_document, c.detected_timestamp "
        "FROM provisions p JOIN change_log c ON c.change_id = p.latest_change_id "
        "WHERE p.provision_id = ? AND c.applied_to_master = 'Y' "
        "AND c.change_origin = 'regulatory' AND c.change_type != 'New Provision'",
        (provision_id,),
    ).fetchone()
    return row


def build_document(conn, spec):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(10.5)

    title = doc.add_heading(spec["title"], level=0)
    title.runs[0].font.color.rgb = HEADING_COLOR

    sub = doc.add_paragraph(spec["subtitle"])
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = GREY
    sub.runs[0].italic = True

    note = doc.add_paragraph(
        "This document is generated from the project database — do not hand-edit; "
        "re-run src/export_word.py after any change is approved."
    )
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = GREY

    doc.add_page_break()

    toc_heading = doc.add_heading("Table of Contents", level=1)
    toc_heading.runs[0].font.color.rgb = HEADING_COLOR

    provisions = fetch_provisions(conn, spec["path"])
    warnings = []

    bookmark_id = 1
    for row in provisions:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_internal_hyperlink(p, f"{row['reference']} — {row['topic_category']}", row["full_text_anchor"])

    doc.add_page_break()

    for row in provisions:
        heading = doc.add_heading(row["reference"], level=2)
        heading.runs[0].font.color.rgb = HEADING_COLOR
        add_bookmark(heading, row["full_text_anchor"], bookmark_id)
        bookmark_id += 1

        meta = doc.add_paragraph()
        meta_run = meta.add_run(
            f"{row['topic_category']}  |  Status: {row['status']}  |  Effective: {row['effective_date']}"
        )
        meta_run.font.size = Pt(8.5)
        meta_run.font.color.rgb = GREY
        meta_run.italic = True
        meta.paragraph_format.space_after = Pt(10)

        amendment = fetch_latest_amendment(conn, row["provision_id"])
        warned = render_provision_body(doc, row["full_text"], amendment)
        if warned:
            warnings.append(row["provision_id"])

        if row["notes"]:
            note_p = doc.add_paragraph()
            note_run = note_p.add_run(f"Note: {row['notes']}")
            note_run.italic = True
            note_run.font.size = Pt(8.5)
            note_run.font.color.rgb = GREY

        doc.add_paragraph()  # spacing between provisions

    return doc, len(provisions), warnings


def main():
    conn = get_connection()
    for spec in DOCX_SPECS:
        doc, count, warnings = build_document(conn, spec)
        out_path = PROJECT_ROOT / spec["path"]
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)
        print(f"Wrote {count} provisions -> {out_path}")
        if warnings:
            print(f"  {len(warnings)} warning(s), see stderr: {', '.join(warnings)}")
    conn.close()


if __name__ == "__main__":
    main()
