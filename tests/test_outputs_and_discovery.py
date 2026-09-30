"""
Phase 3 of the 30 September 2026 audit fixes — what the generated documents say,
and what the discovery layer does when a page misbehaves.

H-4  All 79 "Open full text" links in the distributed Excel file pointed into a
     deleted sandbox folder.
M-1  Only the most recent change to a provision was ever shown.
H-6  A data correction landing after a real amendment erased that amendment's
     highlighting, caption and Excel cell — while the README said data
     corrections "never affect rendering".
M-5  Source_Log claimed to show "what is being watched right now" and held
     7 rows for 4 watched URLs.
M-6  The discovery scrape never checked the count the page itself reports.
M-7  One malformed Gazette ID threw away the whole day's listing.
M-12 The discovery window is two months, so an outage longer than a month lost
     documents permanently.
L-3  The In_Force formula showed #VALUE! for an unreadable date.
L-8  Neither Word document said when it was made, what it was based on, or that
     it is not legal advice.

No network. No AI calls.
"""
from __future__ import annotations

import os
import re
import sys
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from docx.enum.text import WD_COLOR_INDEX

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import discover_documents as dd
import export_excel as ee
import export_word as ew

from conftest import seed_provision, seed_source_log

RULES_SPEC = ew.DOCX_SPECS[0]


def _insert_change(conn, *, change_id, provision_id, change_origin, old_full_text,
                    new_full_text, detected_timestamp, change_type="Correction",
                    source_document="G.S.R. 892(E)", point_latest=True):
    conn.execute(
        """INSERT INTO change_log
           (change_id, detected_timestamp, provision_id, change_type, change_origin,
            old_value_summary, new_value_summary, old_full_text, new_full_text,
            source_document, source_url, detected_by, confidence_score,
            review_status, reviewed_by, review_date, applied_to_master, notes)
           VALUES (?, ?, ?, ?, ?, 'old', 'new', ?, ?, ?, 'url', 'test', 0.9,
                   'Approved', 'auto', ?, 'Y', 'test')""",
        (change_id, detected_timestamp, provision_id, change_type, change_origin,
         old_full_text, new_full_text, source_document, detected_timestamp[:10]),
    )
    if point_latest:
        conn.execute("UPDATE provisions SET latest_change_id = ? WHERE provision_id = ?",
                     (change_id, provision_id))
    conn.commit()


def _highlights(doc):
    return [run.text for p in doc.paragraphs for run in p.runs
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]


def _captions(doc):
    return [run.text for p in doc.paragraphs for run in p.runs if "Amended by" in run.text]


# --------------------------------------------------------------------------
# H-4: the Excel hyperlinks must work for whoever opens the file
# --------------------------------------------------------------------------

def test_hyperlinks_are_relative_to_the_workbook():
    target = ee._workbook_relative("docs/DPDP_Act_2023.docx")
    assert target == "../docs/DPDP_Act_2023.docx"


def test_no_hyperlink_is_an_absolute_path(conn):
    """The exact check the audit asked for: no target may start with '/' or a
    drive letter. Every one of the 79 links in the committed workbook did."""
    seed_provision(conn, provision_id="DPDPR-R23", full_text="text",
                    full_text_path="docs/DPDP_Rules_2025.docx", full_text_anchor="DPDPR_R23")
    seed_provision(conn, provision_id="DPDPA-S1", full_text="text", sort_order=2,
                    full_text_path="docs/DPDP_Act_2023.docx", full_text_anchor="DPDPA_S1",
                    instrument_type="Act Section", reference="Section 1")
    provisions = conn.execute(f"SELECT {','.join(ee.MP_COLS)} FROM provisions").fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ee.write_rows(ws, ee.MP_COLS, provisions,
                  full_text_col="full_text_path", anchor_col="full_text_anchor")

    targets = [
        ws.cell(row=r, column=ee.MP_COL_IDX["full_text_path"]).hyperlink.target
        for r in range(2, ws.max_row + 1)
    ]
    assert targets, "no hyperlinks were written at all"
    for target in targets:
        assert not target.startswith("/"), target
        assert not re.match(r"^[A-Za-z]:", target), target
        assert target.startswith("../docs/"), target
        assert "#" in target, "the bookmark anchor must be kept"


def test_the_link_really_points_at_the_document(tmp_path):
    """Resolve the relative link the way Excel would — against the workbook's
    own folder — and check the file is there."""
    target = ee._workbook_relative("docs/DPDP_Rules_2025.docx")
    resolved = (ee.OUT_PATH.parent / target).resolve()
    assert resolved == (ee.PROJECT_ROOT / "docs" / "DPDP_Rules_2025.docx").resolve()
    assert resolved.exists(), f"{resolved} does not exist"


# --------------------------------------------------------------------------
# H-6: a data correction must not erase an earlier amendment
# --------------------------------------------------------------------------

def test_a_later_data_correction_does_not_hide_a_government_amendment(conn):
    """
    The audit's S18 scenario. A real amendment, then one of our own data
    corrections, which moves provisions.latest_change_id. The amendment must
    still be highlighted, captioned, and named in the Excel column.
    """
    seed_provision(conn, provision_id="DPDPR-R23", full_text="given in such order.",
                    full_text_anchor="DPDPR_R23")
    _insert_change(conn, change_id="CHG-0001", provision_id="DPDPR-R23",
                    change_origin="regulatory", old_full_text="given in such",
                    new_full_text="given in such order",
                    detected_timestamp="2025-12-11T00:00:00")
    # ...and now a data correction of our own, pointed at by latest_change_id.
    _insert_change(conn, change_id="CHG-0002", provision_id="DPDPR-R23",
                    change_origin="data_correction", old_full_text="given in such order.",
                    new_full_text="given in such order.",
                    detected_timestamp="2026-09-23T00:00:00",
                    source_document="our own fix", point_latest=True)

    assert conn.execute(
        "SELECT latest_change_id FROM provisions WHERE provision_id='DPDPR-R23'"
    ).fetchone()["latest_change_id"] == "CHG-0002"

    doc, _count, warnings = ew.build_document(conn, RULES_SPEC)
    assert _highlights(doc) == ["order"], _highlights(doc)
    assert any("CHG-0001" in c for c in _captions(doc)), _captions(doc)
    assert warnings == []

    cells = ee.fetch_last_regulatory_changes(conn)
    assert cells["DPDPR-R23"][0] == "CHG-0001"


def test_a_data_correction_on_its_own_still_renders_nothing(conn):
    """The rule that was always intended: our own fixes are never highlighted."""
    seed_provision(conn, provision_id="DPDPR-R23", full_text="corrected text.",
                    full_text_anchor="DPDPR_R23")
    _insert_change(conn, change_id="CHG-0001", provision_id="DPDPR-R23",
                    change_origin="data_correction", old_full_text="wrong text",
                    new_full_text="corrected text.",
                    detected_timestamp="2026-09-23T00:00:00")
    doc, _count, warnings = ew.build_document(conn, RULES_SPEC)
    assert _highlights(doc) == []
    assert _captions(doc) == []
    assert warnings == []
    assert ee.fetch_last_regulatory_changes(conn) == {}


# --------------------------------------------------------------------------
# M-1: every government change is shown, not just the newest
# --------------------------------------------------------------------------

def test_two_amendments_to_one_provision_are_both_rendered(conn):
    """
    The audit's S17 scenario. Rule 5 amended in 2025 and again in 2026: both
    changed phrases must be highlighted and both captioned. Before 30 Sep 2026
    the first one vanished.
    """
    seed_provision(
        conn, provision_id="DPDPR-R5", reference="Rule 5",
        full_text="(1) The period shall be twelve months.\n\n(2) The fee shall be five hundred rupees.",
        full_text_anchor="DPDPR_R5",
    )
    _insert_change(conn, change_id="CHG-0001", provision_id="DPDPR-R5",
                    change_origin="regulatory", old_full_text="eighteen months",
                    new_full_text="twelve months",
                    detected_timestamp="2025-12-11T00:00:00",
                    source_document="G.S.R. 892(E)", point_latest=False)
    _insert_change(conn, change_id="CHG-0002", provision_id="DPDPR-R5",
                    change_origin="regulatory", old_full_text="two hundred rupees",
                    new_full_text="five hundred rupees",
                    detected_timestamp="2026-06-01T00:00:00",
                    source_document="G.S.R. 401(E)", point_latest=True)

    doc, _count, warnings = ew.build_document(conn, RULES_SPEC)
    assert warnings == []
    # Only the words that actually changed are highlighted, per change:
    # "eighteen months" -> "twelve months" shares one word ("months"), which is
    # not enough to trust as unchanged, so the whole phrase lights up;
    # "two hundred rupees" -> "five hundred rupees" shares two ("hundred
    # rupees"), so only "five" does. See src/word_diff.py.
    assert sorted(_highlights(doc)) == sorted(["twelve months", "five"]), _highlights(doc)
    captions = _captions(doc)
    assert any("CHG-0001" in c for c in captions), captions
    assert any("CHG-0002" in c for c in captions), captions


def test_the_excel_column_says_how_much_earlier_history_there_is(conn):
    seed_provision(conn, provision_id="DPDPR-R5", reference="Rule 5",
                    full_text="twelve months, five hundred rupees.", full_text_anchor="DPDPR_R5")
    _insert_change(conn, change_id="CHG-0001", provision_id="DPDPR-R5",
                    change_origin="regulatory", old_full_text="a", new_full_text="twelve months",
                    detected_timestamp="2025-12-11T00:00:00", point_latest=False)
    _insert_change(conn, change_id="CHG-0002", provision_id="DPDPR-R5",
                    change_origin="regulatory", old_full_text="b",
                    new_full_text="five hundred rupees",
                    detected_timestamp="2026-06-01T00:00:00")

    provisions = conn.execute(f"SELECT {','.join(ee.MP_COLS)} FROM provisions").fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active
    ee.add_last_regulatory_change_column(ws, provisions, ee.fetch_last_regulatory_changes(conn))
    value = ws.cell(row=2, column=ee.LAST_REG_CHANGE_COL_IDX).value
    assert "CHG-0002" in value
    assert "(+1 earlier)" in value, value


def test_an_amendment_whose_text_has_moved_on_is_captioned_but_not_highlighted(conn):
    """Unchanged behaviour, kept: never mis-highlight. The change is still named
    so the audit trail is visible."""
    seed_provision(conn, provision_id="DPDPR-R23", full_text="a completely different sentence.",
                    full_text_anchor="DPDPR_R23")
    _insert_change(conn, change_id="CHG-0001", provision_id="DPDPR-R23",
                    change_origin="regulatory", old_full_text="given in such",
                    new_full_text="given in such order",
                    detected_timestamp="2025-12-11T00:00:00")
    doc, _count, warnings = ew.build_document(conn, RULES_SPEC)
    assert _highlights(doc) == []
    assert any("CHG-0001" in c for c in _captions(doc))
    assert warnings == ["DPDPR-R23"]


# --------------------------------------------------------------------------
# M-5: Source_Log shows what is being watched right now
# --------------------------------------------------------------------------

def test_the_source_log_sheet_leaves_out_sources_no_longer_watched(conn):
    seed_source_log(conn, document_id="SRC-0001", url="https://example.test/watched.pdf")
    seed_source_log(conn, document_id="SRC-0002", url="https://example.test/retired.pdf")
    conn.execute("UPDATE source_log SET watched = 0 WHERE document_id = 'SRC-0002'")
    conn.commit()

    watched = conn.execute(
        f"SELECT {','.join(ee.SL_COLS)} FROM source_log WHERE watched = 1"
    ).fetchall()
    assert [r["document_id"] for r in watched] == ["SRC-0001"]
    # The retired row is still there for its history — just not shown.
    assert conn.execute("SELECT COUNT(*) c FROM source_log").fetchone()["c"] == 2


def test_watched_defaults_to_one_for_an_existing_row(conn):
    """The migration is additive: an old row keeps being watched unless somebody
    deliberately retires it."""
    seed_source_log(conn, document_id="SRC-0009", url="https://example.test/x.pdf")
    assert conn.execute(
        "SELECT watched FROM source_log WHERE document_id='SRC-0009'"
    ).fetchone()["watched"] == 1


# --------------------------------------------------------------------------
# L-3: the In_Force formula
# --------------------------------------------------------------------------

def test_the_in_force_formula_cannot_show_an_error_value():
    wb = openpyxl.Workbook()
    ws = wb.active
    ee.add_in_force_column(ws, 3)
    formula = ws.cell(row=2, column=ee.IN_FORCE_COL_IDX).value
    assert "IFERROR" in formula, formula
    assert "DATEVALUE" in formula


# --------------------------------------------------------------------------
# L-8: the documents say where they came from
# --------------------------------------------------------------------------

def test_both_documents_carry_provenance_and_a_disclaimer(conn):
    seed_provision(conn, provision_id="DPDPR-R23", full_text="text",
                    full_text_anchor="DPDPR_R23")
    for spec in ew.DOCX_SPECS:
        doc, _count, _warnings = ew.build_document(conn, spec)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Last regenerated:" in text
        assert date.today().isoformat() in text
        assert "not legal advice" in text
        assert "Verify any clause you quote against the official Gazette" in text


def test_the_rules_document_names_the_corrigendum(conn):
    seed_provision(conn, provision_id="DPDPR-R23", full_text="text",
                    full_text_anchor="DPDPR_R23")
    doc, _count, _warnings = ew.build_document(conn, ew.DOCX_SPECS[0])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "G.S.R. 892(E)" in text
    assert "10 December 2025" in text


def test_the_documents_say_how_current_the_text_is(conn):
    seed_provision(conn, provision_id="DPDPR-R23", full_text="given in such order.",
                    full_text_anchor="DPDPR_R23", last_updated_date="2026-09-23")
    _insert_change(conn, change_id="CHG-0001", provision_id="DPDPR-R23",
                    change_origin="regulatory", old_full_text="given in such",
                    new_full_text="given in such order",
                    detected_timestamp="2025-12-11T00:00:00")
    lines = ew.provenance_lines(conn, ew.DOCX_SPECS[0])
    joined = "\n".join(lines)
    assert "Text accurate as at: 2026-09-23" in joined
    assert "Most recent government change recorded here: 2025-12-11" in joined


# --------------------------------------------------------------------------
# M-6 / M-7: the discovery scrape checks itself
# --------------------------------------------------------------------------

class FakePage:
    """The few Playwright calls _egazette_search_month makes, and nothing else."""

    def __init__(self, rows, result_label="Total No. of Gazettes : {n}"):
        self.rows = rows
        self.result_label = result_label

    def goto(self, *a, **k):
        pass

    def route(self, *a, **k):
        pass

    def click(self, *a, **k):
        pass

    def wait_for_selector(self, *a, **k):
        pass

    def wait_for_timeout(self, *a, **k):
        pass

    def select_option(self, *a, **k):
        pass

    def expect_navigation(self, *a, **k):
        class _Ctx:
            def __enter__(self_inner):
                return None

            def __exit__(self_inner, *exc):
                return False
        return _Ctx()

    def inner_text(self, selector):
        if self.result_label is None:
            raise RuntimeError("label missing")
        return self.result_label.format(n=len(self.rows))

    def eval_on_selector_all(self, *a, **k):
        return self.rows


def _row(gazette_id, subject="A notification", issue_date="11-Dec-2025"):
    return {"subject": subject, "issueDate": issue_date, "gazetteId": gazette_id}


def test_the_reported_count_is_read_from_the_page():
    rows = [_row("CG-DL-E-12122025-268455"), _row("CG-DL-E-12122025-268456")]
    got_rows, reported = dd._egazette_search_month(FakePage(rows), 2025, "December")
    assert got_rows == rows
    assert reported == 2


def _fake_playwright(monkeypatch, rows, result_label="Total No. of Gazettes : {n}"):
    """Point fetch_egazette_meity at a fake browser that serves `rows`."""
    page = FakePage(rows, result_label=result_label)

    class FakeBrowser:
        def new_page(self, **kwargs):
            return page

        def close(self):
            pass

    class FakePlaywright:
        chromium = type("C", (), {"launch": staticmethod(lambda **k: FakeBrowser())})()

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    import playwright.sync_api as sync_api
    monkeypatch.setattr(sync_api, "sync_playwright", lambda: FakePlaywright())
    monkeypatch.setattr(dd, "_months_to_search", lambda *a, **k: [(2025, 12, "December")])
    return page


def test_a_month_whose_count_disagrees_is_a_loud_failure(conn, monkeypatch):
    """
    M-6. If the page says it found 5 and only 2 rows could be read, the listing
    is not the whole month — and a short listing looks exactly like "nothing
    new", which is the failure that matters.
    """
    rows = [_row("CG-DL-E-12122025-268455"), _row("CG-DL-E-12122025-268456")]
    _fake_playwright(monkeypatch, rows, result_label="Total No. of Gazettes : 5")
    with pytest.raises(RuntimeError, match="said it found 5"):
        dd.fetch_egazette_meity(conn=conn, today=date(2025, 12, 15))


def test_a_month_whose_count_agrees_is_accepted(conn, monkeypatch):
    rows = [_row("CG-DL-E-12122025-268455"), _row("CG-DL-E-12122025-268456")]
    _fake_playwright(monkeypatch, rows)
    result = dd.fetch_egazette_meity(conn=conn, today=date(2025, 12, 15))
    assert len(result["items"]) == 2
    assert result["warnings"] == []


def test_an_unreadable_count_label_warns_but_keeps_going(conn, monkeypatch):
    rows = [_row("CG-DL-E-12122025-268455")]
    _fake_playwright(monkeypatch, rows, result_label=None)
    result = dd.fetch_egazette_meity(conn=conn, today=date(2025, 12, 15))
    assert len(result["items"]) == 1
    assert any("could not be read" in w for w in result["warnings"]), result["warnings"]


def test_an_unreadable_count_label_is_a_warning_not_a_crash():
    rows = [_row("CG-DL-E-12122025-268455")]
    got_rows, reported = dd._egazette_search_month(FakePage(rows, result_label=None),
                                                   2025, "December")
    assert got_rows == rows
    assert reported is None


def test_a_bad_gazette_id_keeps_the_good_rows(conn, monkeypatch):
    """
    M-7. One malformed ID used to raise and throw away the whole day's listing.
    Now the good rows survive and the bad one is named for a human.
    """
    rows = [
        _row("CG-DL-E-12122025-268455", subject="Good one"),
        _row("NOT-A-GAZETTE-ID", subject="Odd one"),
        _row("CG-DL-E-12122025-268457", subject="Another good one"),
    ]
    _fake_playwright(monkeypatch, rows)
    result = dd.fetch_egazette_meity(conn=conn, today=date(2025, 12, 15))
    urls = sorted(item["url"] for item in result["items"])
    assert urls == [
        "https://egazette.gov.in/WriteReadData/2025/268455.pdf",
        "https://egazette.gov.in/WriteReadData/2025/268457.pdf",
    ], urls
    assert any("NOT-A-GAZETTE-ID" in w for w in result["warnings"]), result["warnings"]
    assert any("by hand" in w for w in result["warnings"])


def test_a_discovery_warning_reaches_the_runs_error_list(conn):
    src = {
        "name": "src1",
        "all_items_relevant": False,
        "fetch": lambda **kw: {
            "items": [], "bucket_counts": {"src1": 0},
            "warnings": ["a Gazette ID could not be parsed — please look it up by hand"],
        },
    }
    errors: list[str] = []
    with patch.object(dd, "DISCOVERY_SOURCES", [src]):
        dd.discover_all(conn, errors)
    assert any("look it up by hand" in e for e in errors), errors


# --------------------------------------------------------------------------
# M-12: the search window widens after an outage
# --------------------------------------------------------------------------

def test_an_ordinary_day_searches_two_months(conn):
    conn.execute(
        "INSERT INTO discovery_run_log (discovery_source, run_date, items_found) VALUES (?,?,?)",
        ("egazette-meity:__init__", "2026-09-29", 5))
    conn.commit()
    months = dd._months_to_search(conn, date(2026, 9, 30), "egazette-meity")
    assert len(months) == 2
    assert months[0][:2] == (2026, 9)
    assert months[1][:2] == (2026, 8)


def test_a_source_that_has_never_run_searches_two_months(conn):
    months = dd._months_to_search(conn, date(2026, 9, 30), "egazette-meity")
    assert len(months) == 2


def test_a_long_outage_widens_the_window_and_says_so(conn):
    conn.execute(
        "INSERT INTO discovery_run_log (discovery_source, run_date, items_found) VALUES (?,?,?)",
        ("egazette-meity:2026-05", "2026-05-02", 4))
    conn.commit()
    warnings: list[str] = []
    months = dd._months_to_search(conn, date(2026, 9, 30), "egazette-meity", warnings)
    assert [(y, m) for y, m, _name in months] == [
        (2026, 5), (2026, 6), (2026, 7), (2026, 8), (2026, 9)
    ], months
    assert warnings and "searched 5 month(s)" in warnings[0], warnings


def test_a_very_long_outage_is_capped_and_the_gap_is_named(conn):
    conn.execute(
        "INSERT INTO discovery_run_log (discovery_source, run_date, items_found) VALUES (?,?,?)",
        ("egazette-meity:2024-01", "2024-01-05", 4))
    conn.commit()
    warnings: list[str] = []
    months = dd._months_to_search(conn, date(2026, 9, 30), "egazette-meity", warnings)
    assert len(months) == dd.DISCOVERY_MAX_MONTHS
    assert months[-1][:2] == (2026, 9)
    assert any("capped" in w and "by hand" in w for w in warnings), warnings
