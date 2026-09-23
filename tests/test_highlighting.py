"""
Phase 8: D1/D2's highlighting rule is a permanent guard — only a real
regulatory change is ever highlighted, and only the words that changed.
Covers: a data_correction produces 0 highlighted docx runs and 0 yellow
Excel cells; a synthetic regulatory change highlights only the changed
words; the G.S.R. 892(E) corrigendum wording highlights exactly as D2
specifies; the email lists regulatory changes only.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx.enum.text import WD_COLOR_INDEX
import openpyxl

import export_excel as ee
import export_word as ew
from word_diff import change_snippet
from notify import _build_body, _regulatory_changes

from conftest import seed_provision


DOCX_SPEC = {
    "path": "docs/DPDP_Rules_2025.docx",
    "title": "Test",
    "subtitle": "Test",
}


def _insert_change(conn, *, change_id, provision_id, change_type, change_origin,
                    old_full_text, new_full_text, point_latest_change_id=True,
                    source_document="Test Source", detected_timestamp="2026-09-23T00:00:00"):
    conn.execute(
        """INSERT INTO change_log
           (change_id, detected_timestamp, provision_id, change_type, change_origin,
            old_value_summary, new_value_summary, old_full_text, new_full_text,
            source_document, source_url, detected_by, confidence_score,
            review_status, reviewed_by, review_date, applied_to_master, notes)
           VALUES (?, ?, ?, ?, ?, 'old summary', 'new summary', ?, ?, ?, 'url', 'test', 0.9,
                   'Approved', 'auto', ?, 'Y', 'test')""",
        (change_id, detected_timestamp, provision_id, change_type, change_origin,
         old_full_text, new_full_text, source_document, detected_timestamp[:10]),
    )
    if point_latest_change_id:
        conn.execute("UPDATE provisions SET latest_change_id = ? WHERE provision_id = ?",
                     (change_id, provision_id))
    conn.commit()


def _highlighted_runs(doc):
    """(text, style) for every run with a highlight or a strike, across the doc."""
    out = []
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                out.append((run.text, "highlight"))
            elif run.font.strike:
                out.append((run.text, "strike"))
    return out


# --------------------------------------------------------------------------
# data_correction: 0 highlights anywhere
# --------------------------------------------------------------------------

def test_data_correction_produces_zero_docx_highlights(conn):
    seed_provision(conn, full_text="corrected text now on file.")
    _insert_change(
        conn, change_id="CHG-0001", provision_id="DPDPR-R23",
        change_type="Correction", change_origin="data_correction",
        old_full_text="wrong text", new_full_text="corrected text now on file.",
        point_latest_change_id=True,  # even pointed at, data_correction must never render
    )
    doc, count, warnings = ew.build_document(conn, DOCX_SPEC)
    assert _highlighted_runs(doc) == []
    assert warnings == []


def test_data_correction_produces_zero_excel_yellow_cells(conn):
    seed_provision(conn, full_text="corrected text now on file.")
    _insert_change(
        conn, change_id="CHG-0001", provision_id="DPDPR-R23",
        change_type="Correction", change_origin="data_correction",
        old_full_text="wrong text", new_full_text="corrected text now on file.",
    )
    changes = [dict(r) for r in conn.execute(f"SELECT {','.join(ee.CL_COLS)} FROM change_log")]
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, h in enumerate(ee.CL_COLS, start=1):
        ws.cell(row=1, column=i, value=h)
    ee.write_rows(ws, ee.CL_COLS, changes)
    ee.highlight_regulatory_rows(ws, changes)

    yellow_cells = [
        (r, c) for r in range(2, ws.max_row + 1) for c in range(1, len(ee.CL_COLS) + 1)
        if ws.cell(row=r, column=c).fill.start_color.rgb not in (None, "00000000")
    ]
    assert yellow_cells == []


# --------------------------------------------------------------------------
# Regulatory change: only the changed words are highlighted
# --------------------------------------------------------------------------

def test_regulatory_change_highlights_only_the_changed_word(conn):
    seed_provision(conn, full_text="given in such order.")
    _insert_change(
        conn, change_id="CHG-0002", provision_id="DPDPR-R23",
        change_type="Correction", change_origin="regulatory",
        old_full_text="given in such", new_full_text="given in such order",
        source_document="G.S.R. 892(E)",
    )
    doc, count, warnings = ew.build_document(conn, DOCX_SPEC)
    runs = _highlighted_runs(doc)
    assert runs == [("order", "highlight")], f"expected only 'order' highlighted, got {runs}"
    assert warnings == []
    # caption present
    caption_texts = [r.text for p in doc.paragraphs for r in p.runs if "Amended by" in r.text]
    assert any("G.S.R. 892(E)" in t and "CHG-0002" in t for t in caption_texts)


def test_regulatory_change_no_longer_matching_current_text_renders_plainly_with_warning(conn):
    """If a later data_correction changed the surrounding text so the
    regulatory change's new_full_text is no longer found verbatim, render
    plainly with the caption and log a warning — never mis-highlight."""
    seed_provision(conn, full_text="a completely different sentence now.")
    _insert_change(
        conn, change_id="CHG-0003", provision_id="DPDPR-R23",
        change_type="Correction", change_origin="regulatory",
        old_full_text="given in such", new_full_text="given in such order",
    )
    doc, count, warnings = ew.build_document(conn, DOCX_SPEC)
    assert _highlighted_runs(doc) == []
    assert warnings == ["DPDPR-R23"]


# --------------------------------------------------------------------------
# G.S.R. 892(E) corrigendum wording — exactly as D2 specifies
# --------------------------------------------------------------------------

def test_corrigendum_wording_highlights_exactly_as_specified(conn):
    """
    The real DPDPR-R1 has "of this Gazette" -> "in the Official Gazette"
    corrected in BOTH sub-rule (3) and sub-rule (4) — the corrigendum's
    two separate line-level items (page 24 lines 22 and 24) both apply the
    identical phrase fix within the SAME provision row, not two different
    provisions. Sub-rule (2), included here, ALREADY correctly said "in
    the Official Gazette" before this correction and must NOT be
    highlighted — only the two genuinely corrected occurrences may be
    (this is exactly the collision scripts/apply_rules_corrigendum_2026-09-23.py's
    widen_until_unambiguous() exists to avoid).
    """
    seed_provision(
        conn, provision_id="DPDPR-R1",
        full_text=(
            "**Rule 1 — Short title and commencement**\n\n"
            "(2) Rules 1, 2 and 17 to 21 shall come into force on the date "
            "of their publication in the Official Gazette.\n\n"
            "(3) Rule 4 shall come into force one year after the date of "
            "publication in the Official Gazette.\n\n"
            "(4) Rules 3, 5 to 16, 22 and 23 shall come into force eighteen "
            "months after the date of publication in the Official Gazette."
        ),
        full_text_anchor="anchor_r1", sort_order=1,
    )
    _insert_change(
        conn, change_id="CHG-0100", provision_id="DPDPR-R1",
        change_type="Correction", change_origin="regulatory",
        old_full_text="of publication of this Gazette", new_full_text="of publication in the Official Gazette",
        source_document="G.S.R. 892(E)",
    )
    seed_provision(conn, provision_id="DPDPR-SDF-R13", full_text="the Departments concerned.",
                   full_text_anchor="anchor_r13", sort_order=2)
    _insert_change(
        conn, change_id="CHG-0101", provision_id="DPDPR-SDF-R13",
        change_type="Correction", change_origin="regulatory",
        old_full_text="Department", new_full_text="Departments", source_document="G.S.R. 892(E)",
    )
    seed_provision(conn, provision_id="DPDPR-R23", full_text="given in such order.",
                   full_text_anchor="anchor_r23", sort_order=3)
    _insert_change(
        conn, change_id="CHG-0102", provision_id="DPDPR-R23",
        change_type="Correction", change_origin="regulatory",
        old_full_text="given in such", new_full_text="given in such order", source_document="G.S.R. 892(E)",
    )

    doc, count, warnings = ew.build_document(conn, DOCX_SPEC)
    assert warnings == []
    highlighted_texts = [t for t, style in _highlighted_runs(doc) if style == "highlight"]
    assert highlighted_texts == ["in the Official Gazette", "in the Official Gazette", "Departments", "order"], \
        highlighted_texts


# --------------------------------------------------------------------------
# Email: regulatory changes only
# --------------------------------------------------------------------------

def test_email_lists_regulatory_changes_only():
    reg = {"provision_id": "DPDPR-R23", "reference": "Rule 23(1)", "change_type": "Correction",
           "change_origin": "regulatory", "old_full_text": "given in such",
           "new_full_text": "given in such order", "new_value_summary": "x"}
    data_corr = {**reg, "provision_id": "DPDPR-R1", "reference": "Rule 1(1)",
                 "change_origin": "data_correction"}

    reg_only = _regulatory_changes([reg, data_corr])
    assert reg_only == [reg]

    subject, body = _build_body([reg, data_corr], [])
    assert "1 change" in subject
    assert "Rule 23(1)" in body
    assert "Rule 1(1)" not in body


def test_word_diff_snippet_matches_docx_highlighting():
    """The email snippet and the docx highlight are built from the same
    word_diff logic — they must never disagree about what changed."""
    old_snip, new_snip = change_snippet("of this Gazette", "in the Official Gazette")
    assert new_snip == "in the Official Gazette"
    assert old_snip == "of this Gazette"
